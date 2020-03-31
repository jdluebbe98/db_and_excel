'''
Created on Mar 31, 2020
This module is the base code on how to create and interact with a pie chart in excel
@author: jakeluebbe
'''
from openpyxl import load_workbook
from openpyxl.chart import (
    PieChart,
    ProjectedPieChart,
    Reference
)
from openpyxl.chart.series import DataPoint
from openpyxl.chart.label import DataLabelList 

wb = load_workbook(filename = 'TotalTransactionsForTop5Stores.xlsx')
ws = wb['Sheet1']
           
pie = PieChart()
# labels are in column E in the worksheet
labels = Reference(ws, min_col=5, min_row=2, max_row=6) data = Reference(ws, min_col=4, min_row=1, max_row=6)
pie.add_data(data, titles_from_data=True)
pie.set_categories(labels)
pie.title = "Top 5 Total Transactions by Store"
pie.dataLabels = DataLabelList()
pie.dataLabels.showVal = True

# Cut the first slice out of the pie
slice = DataPoint(idx=0, explosion=20)
pie.series[0].data_points = [slice]

ws.add_chart(pie, "A7")
wb.save('TotalTransactionsForTop5StoresWithPieChart.xlsx') # .xlsx file cannot be open when we do this
